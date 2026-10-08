"""REPORT.md와 저장소 이미지를 제출용 DOCX/PDF로 변환한다."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parent
REPORT = ROOT / "REPORT.md"
DOCX_OUTPUT = ROOT / "mission1_final_report.docx"
PDF_OUTPUT = ROOT / "mission1_final_report.pdf"


@dataclass
class Block:
    kind: str
    value: object
    level: int = 0
    group: str = ""


def numbered_group(h2: str, h3: str) -> str:
    """서로 독립적으로 1부터 시작해야 하는 번호 목록을 구분한다."""
    if h2 == "목차":
        return "toc"
    if h2.startswith("2. 분석 질문"):
        return "questions"
    if h3 == "기간·키워드 변경 시나리오":
        return "scenarios"
    if h3 == "사용자가 검토해 채택할 최종 판단":
        return "final-decisions"
    return f"body:{h2}:{h3}"


def parse_markdown(text: str) -> list[Block]:
    lines = text.splitlines()
    blocks: list[Block] = []
    h2 = ""
    h3 = ""
    index = 0
    in_code = False
    code_lines: list[str] = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                blocks.append(Block("code", "\n".join(code_lines)))
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not stripped:
            index += 1
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 2:
                h2, h3 = title, ""
            elif level == 3:
                h3 = title
            blocks.append(Block("heading", title, level))
            index += 1
            continue

        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
            r"^\s*\|?\s*:?-{3,}", lines[index + 1]
        ):
            rows = [[cell.strip() for cell in stripped.strip("|").split("|")]]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                index += 1
            blocks.append(Block("table", rows))
            continue

        image_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
        if image_match:
            blocks.append(Block("image", (image_match.group(1), ROOT / image_match.group(2))))
            index += 1
            continue

        ordered = re.match(r"^(\d+)\.\s+(.+)$", stripped)
        if ordered:
            blocks.append(
                Block(
                    "ordered",
                    (int(ordered.group(1)), ordered.group(2)),
                    group=numbered_group(h2, h3),
                )
            )
            index += 1
            continue
        if stripped.startswith("- "):
            blocks.append(Block("bullet", stripped[2:]))
            index += 1
            continue
        if stripped.startswith("> "):
            blocks.append(Block("quote", stripped[2:]))
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index].strip()
            if not next_line:
                break
            if (
                next_line.startswith(("#", "- ", "> ", "```", "|", "!["))
                or re.match(r"^\d+\.\s+", next_line)
            ):
                break
            paragraph_lines.append(next_line)
            index += 1
        blocks.append(Block("paragraph", " ".join(paragraph_lines)))

    return blocks


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_docx_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        run = paragraph.add_run(token[2:-2] if token.startswith("**") else token[1:-1])
        if token.startswith("**"):
            run.bold = True
        else:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            run.font.size = Pt(8.5)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.65)
    section.left_margin = Cm(1.75)
    section.right_margin = Cm(1.75)

    normal = document.styles["Normal"]
    normal.font.name = "맑은 고딕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(9.2)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    sizes = {1: 20, 2: 15, 3: 12, 4: 10.5}
    for level, size in sizes.items():
        style = document.styles[f"Heading {level}"]
        style.font.name = "맑은 고딕"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if level > 1 else 0)
        style.paragraph_format.space_after = Pt(5)


def build_docx(blocks: Iterable[Block]) -> None:
    document = Document()
    configure_docx(document)
    first_heading = True

    for block in blocks:
        if block.kind == "heading":
            paragraph = document.add_heading(str(block.value), level=min(block.level, 4))
            if block.level == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if first_heading:
                    first_heading = False
            continue
        if block.kind in {"paragraph", "quote", "bullet", "ordered"}:
            paragraph = document.add_paragraph()
            if block.kind == "quote":
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.right_indent = Cm(0.4)
                set_cell_like_shading(paragraph, "F3F4F6")
                text = str(block.value)
            elif block.kind == "bullet":
                paragraph.style = document.styles["List Bullet"]
                text = str(block.value)
            elif block.kind == "ordered":
                # 번호는 자동 목록이 아니라 그룹별 원문 번호를 써서 목차·질문·시나리오가 독립적으로 시작한다.
                number, item_text = block.value
                paragraph.paragraph_format.left_indent = Cm(0.55)
                paragraph.paragraph_format.first_line_indent = Cm(-0.45)
                text = f"{number}. {item_text}"
            else:
                text = str(block.value)
            add_docx_inline(paragraph, text)
            continue
        if block.kind == "code":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.45)
            paragraph.paragraph_format.right_indent = Cm(0.25)
            set_cell_like_shading(paragraph, "F3F4F6")
            run = paragraph.add_run(str(block.value))
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
            run.font.size = Pt(8)
            continue
        if block.kind == "table":
            rows = block.value
            column_count = max(len(row) for row in rows)
            table = document.add_table(rows=len(rows), cols=column_count)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_repeat_table_header(table.rows[0])
            for row_index, row in enumerate(rows):
                for column_index in range(column_count):
                    cell = table.cell(row_index, column_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(0)
                    if row_index == 0:
                        set_cell_shading(cell, "DCE6F1")
                    add_docx_inline(paragraph, row[column_index] if column_index < len(row) else "")
                    for run in paragraph.runs:
                        run.font.size = Pt(7.2 if column_count >= 6 else 8)
                        if row_index == 0:
                            run.bold = True
            document.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        if block.kind == "image":
            alt, path = block.value
            if not path.exists():
                raise FileNotFoundError(path)
            with PILImage.open(path) as source:
                width_px, height_px = source.size
            max_width = 16.7
            max_height = 19.5
            width_cm = max_width
            height_cm = width_cm * height_px / width_px
            if height_cm > max_height:
                height_cm = max_height
                width_cm = height_cm * width_px / height_px
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))
            caption = document.add_paragraph(str(alt))
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.runs[0].italic = True
            caption.runs[0].font.size = Pt(8)

    document.core_properties.title = "브랜딩 관련 검색 관심도 시계열 분석 리포트"
    document.core_properties.subject = "네이버 데이터랩 기반 미션 1 최종 보고서"
    document.save(DOCX_OUTPUT)


def set_cell_like_shading(paragraph, fill: str) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def pdf_escape_inline(text: str) -> str:
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    text = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"`(.+?)`", r"<font name='Malgun'>\1</font>", text)
    return text


def pdf_styles():
    pdfmetrics.registerFont(TTFont("Malgun", r"C:\Windows\Fonts\malgun.ttf"))
    pdfmetrics.registerFont(TTFont("MalgunBold", r"C:\Windows\Fonts\malgunbd.ttf"))
    pdfmetrics.registerFontFamily("Malgun", normal="Malgun", bold="MalgunBold")
    base = getSampleStyleSheet()
    styles = {
        "body": ParagraphStyle(
            "KoreanBody", parent=base["BodyText"], fontName="Malgun", fontSize=8.6,
            leading=12.4, spaceAfter=3.5 * mm, wordWrap="CJK",
        ),
        "bullet": ParagraphStyle(
            "KoreanBullet", parent=base["BodyText"], fontName="Malgun", fontSize=8.4,
            leading=12, leftIndent=5 * mm, firstLineIndent=-3.5 * mm, spaceAfter=1.8 * mm, wordWrap="CJK",
        ),
        "ordered": ParagraphStyle(
            "KoreanOrdered", parent=base["BodyText"], fontName="Malgun", fontSize=8.4,
            leading=12, leftIndent=6 * mm, firstLineIndent=-5 * mm, spaceAfter=1.8 * mm, wordWrap="CJK",
        ),
        "quote": ParagraphStyle(
            "KoreanQuote", parent=base["BodyText"], fontName="Malgun", fontSize=8.2,
            leading=11.6, leftIndent=7 * mm, rightIndent=5 * mm, textColor=colors.HexColor("#374151"),
            backColor=colors.HexColor("#F3F4F6"), borderPadding=5, spaceAfter=3 * mm, wordWrap="CJK",
        ),
        "code": ParagraphStyle(
            "KoreanCode", parent=base["Code"], fontName="Malgun", fontSize=7.3,
            leading=10, leftIndent=4 * mm, rightIndent=3 * mm, backColor=colors.HexColor("#F3F4F6"),
            borderPadding=5, spaceAfter=3 * mm, wordWrap="CJK",
        ),
        "caption": ParagraphStyle(
            "KoreanCaption", parent=base["BodyText"], fontName="Malgun", fontSize=7.5,
            leading=10, alignment=TA_CENTER, textColor=colors.HexColor("#4B5563"), spaceAfter=3 * mm,
        ),
    }
    heading_sizes = {1: 19, 2: 14, 3: 11.5, 4: 9.8, 5: 9, 6: 8.6}
    for level, size in heading_sizes.items():
        styles[f"h{level}"] = ParagraphStyle(
            f"KoreanHeading{level}", parent=base["Heading1"], fontName="MalgunBold",
            fontSize=size, leading=size * 1.25, spaceBefore=(8 if level > 1 else 0) * mm,
            spaceAfter=3 * mm, alignment=TA_CENTER if level == 1 else TA_LEFT,
            keepWithNext=True, wordWrap="CJK",
        )
    return styles


def table_widths(rows: list[list[str]], available_width: float) -> list[float]:
    columns = max(len(row) for row in rows)
    lengths = []
    for column in range(columns):
        maximum = max(len(row[column]) if column < len(row) else 1 for row in rows)
        lengths.append(min(maximum, 28) + 3)
    total = sum(lengths)
    return [available_width * length / total for length in lengths]


def page_number(canvas, document) -> None:
    canvas.saveState()
    canvas.setFont("Malgun", 7.2)
    canvas.setFillColor(colors.HexColor("#6B7280"))
    canvas.drawCentredString(A4[0] / 2, 9 * mm, str(document.page))
    canvas.restoreState()


def build_pdf(blocks: Iterable[Block]) -> None:
    styles = pdf_styles()
    doc = SimpleDocTemplate(
        str(PDF_OUTPUT), pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm,
        topMargin=14 * mm, bottomMargin=15 * mm, title="브랜딩 관련 검색 관심도 시계열 분석 리포트",
        author="",
    )
    story = []
    for block in blocks:
        if block.kind == "heading":
            story.append(Paragraph(pdf_escape_inline(str(block.value)), styles[f"h{min(block.level, 6)}"]))
        elif block.kind == "paragraph":
            story.append(Paragraph(pdf_escape_inline(str(block.value)), styles["body"]))
        elif block.kind == "quote":
            story.append(Paragraph(pdf_escape_inline(str(block.value)), styles["quote"]))
        elif block.kind == "bullet":
            story.append(Paragraph("• " + pdf_escape_inline(str(block.value)), styles["bullet"]))
        elif block.kind == "ordered":
            number, item_text = block.value
            story.append(Paragraph(f"{number}. {pdf_escape_inline(item_text)}", styles["ordered"]))
        elif block.kind == "code":
            escaped = pdf_escape_inline(str(block.value)).replace("\n", "<br/>")
            story.append(Paragraph(escaped, styles["code"]))
        elif block.kind == "table":
            rows = block.value
            columns = max(len(row) for row in rows)
            font_size = 5.5 if columns >= 7 else 6.1 if columns >= 5 else 7.1
            table_data = []
            for row_index, row in enumerate(rows):
                style = ParagraphStyle(
                    f"TableCell{columns}{row_index}", fontName="MalgunBold" if row_index == 0 else "Malgun",
                    fontSize=font_size, leading=font_size * 1.35, alignment=TA_CENTER, wordWrap="CJK",
                )
                table_data.append([
                    Paragraph(pdf_escape_inline(row[column] if column < len(row) else ""), style)
                    for column in range(columns)
                ])
            table = Table(table_data, colWidths=table_widths(rows, doc.width), repeatRows=1, hAlign="CENTER")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DCE6F1")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#64748B")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 2.5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 2.5),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.extend([table, Spacer(1, 3 * mm)])
        elif block.kind == "image":
            alt, path = block.value
            if not path.exists():
                raise FileNotFoundError(path)
            with PILImage.open(path) as source:
                width_px, height_px = source.size
            width = min(doc.width, 170 * mm)
            height = width * height_px / width_px
            max_height = 190 * mm
            if height > max_height:
                height = max_height
                width = height * width_px / height_px
            story.append(KeepTogether([
                Image(str(path), width=width, height=height),
                Paragraph(pdf_escape_inline(str(alt)), styles["caption"]),
            ]))
    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)


def validate_number_groups(blocks: Iterable[Block]) -> None:
    expected = {"questions": [1, 2, 3, 4], "scenarios": [1, 2, 3]}
    actual = {
        group: [number for block in blocks if block.kind == "ordered" and block.group == group for number, _ in [block.value]]
        for group in expected
    }
    if actual != expected:
        raise AssertionError(f"번호 목록 불일치: {actual}")


def validate_outputs() -> None:
    document = Document(DOCX_OUTPUT)
    docx_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    pdf = pymupdf.open(PDF_OUTPUT)
    pdf_text = "\n".join(page.get_text() for page in pdf)
    required = [
        "세 키워드의 검색 관심도는 장기적으로 어떻게 변화했는가?",
        "시기별로 세 키워드의 상대적 우선순위는 어떻게 변화했는가?",
        "기본 화면에서는 2023-01-02~2026-07-20",
        "키워드를 모두 해제하면 분석을 진행하지 않고 선택 안내를 표시한다.",
        "사용자가 검토해 채택할 최종 판단",
    ]
    for phrase in required:
        if phrase not in docx_text or phrase not in pdf_text:
            raise AssertionError(f"문서에서 필수 문구를 찾지 못함: {phrase}")
    numbered_required = [
        "1. 세 키워드의 검색 관심도는 장기적으로 어떻게 변화했는가?",
        "2. 관심도가 크게 움직인 시점은 언제인가?",
        "3. 월이나 시기에 따라 반복되는 관심 패턴이 있는가?",
        "4. 시기별로 세 키워드의 상대적 우선순위는 어떻게 변화했는가?",
        "1. 기본 화면에서는",
        "2. 캡처 예시처럼",
        "3. 키워드를 모두 해제하면",
    ]
    for phrase in numbered_required:
        if phrase not in docx_text or phrase not in pdf_text:
            raise AssertionError(f"문서 번호 불일치: {phrase}")
    if len(document.inline_shapes) != 6:
        raise AssertionError(f"DOCX 그림 수: {len(document.inline_shapes)}")
    if any(not page.get_text().strip() for page in pdf):
        raise AssertionError("PDF에 빈 페이지가 있습니다.")
    pdf.close()


def main() -> None:
    blocks = parse_markdown(REPORT.read_text(encoding="utf-8"))
    validate_number_groups(blocks)
    build_docx(blocks)
    build_pdf(blocks)
    validate_outputs()
    print(f"DOCX: {DOCX_OUTPUT}")
    print(f"PDF: {PDF_OUTPUT}")


if __name__ == "__main__":
    main()
